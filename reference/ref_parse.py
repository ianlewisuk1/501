"""Reference extractor for one team from an RDL 'Tons' newsletter workbook.
Used only to generate golden JSON fixtures for the TS build. Not production code."""
import json, re, sys
import openpyxl
from openpyxl.utils import column_index_from_string as ci

TEAM = "Area 501"
DIV = "C"
ALIASES = {"area 501", "area"}  # structured cells only; prose uses "Area 501"


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def is_us(s):
    return norm(s) in ALIASES


def v(ws, col, row):
    return ws.cell(row=row, column=ci(col)).value


def dash(x):
    """The newsletter uses '-' for 'nothing here'."""
    return None if x in (None, "", "-") else x


def num(x):
    return x if isinstance(x, (int, float)) else None


def first_last(last_first):
    s = re.sub(r"\s*&\s*$", "", str(last_first)).strip()
    if "," in s:
        last, first = [p.strip() for p in s.split(",", 1)]
        return f"{first} {last}"
    return s


def sheets(wb):
    m = re.match(r"(\w+?)-Wk(\d+)-Pg", wb.sheetnames[0])
    season, week = m.group(1), int(m.group(2))
    pre = f"{season}-Wk{week}-"
    return season, week, lambda pg: wb[pre + pg]


def issue_info(ws):
    s = str(v(ws, "B", 2))
    m = re.search(r"week #(\d+), \w+, (\w+ \d+, \d{4})\s+Issue #(\w+)", s)
    from datetime import datetime
    return datetime.strptime(m.group(2), "%B %d, %Y").date().isoformat(), m.group(3)


def find_row(ws, col, pred, start=1):
    for r in range(start, ws.max_row + 1):
        if pred(v(ws, col, r)):
            return r
    return None


def page2(ws):
    title = str(v(ws, "A", 1))
    m = re.search(r"WEEK (\d+) \((\d+)/(\d+)/(\d+)\)", title)
    res_week = int(m.group(1))
    res_date = f"20{m.group(4)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    top = find_row(ws, "A", lambda x: norm(x) == f'"{DIV.lower()}" division')
    standings, results_rows = [], []
    r = top
    while r <= ws.max_row and v(ws, "C", r) not in (None, ""):
        standings.append({
            "rank": v(ws, "C", r), "team": str(v(ws, "D", r)).strip(),
            "wins": v(ws, "F", r), "losses": v(ws, "G", r), "points": v(ws, "H", r),
        })
        results_rows.append(r)
        r += 1
    # results: consecutive non-BYE rows pair up; home team is listed second
    rows = [x for x in results_rows if norm(v(ws, "M", x)) != "(bye)"]
    last = None
    for a, b in zip(rows[0::2], rows[1::2]):
        for us, them, home in ((a, b, False), (b, a, True)):
            if is_us(v(ws, "K", us)):
                last = {
                    "week": res_week, "date": res_date,
                    "opponentShort": str(v(ws, "K", them)).strip(), "home": home,
                    "score": {"us": v(ws, "L", us), "them": v(ws, "L", them)},
                    "gamesWon": {k: v(ws, c, us) for k, c in
                                 (("singles301", "M"), ("singlesCricket", "N"), ("doublesCricket", "O"), ("doubles501", "P"))},
                    "tiebreaker1001": dash(v(ws, "Q", us)), "allStarPoints": v(ws, "R", us),
                    "shortHanded": dash(v(ws, "S", us)), "penalties": dash(v(ws, "T", us)),
                }
    if last is None and any(is_us(v(ws, "K", x)) for x in results_rows):
        last = {"week": res_week, "date": res_date, "bye": True}
    return standings, last


def page5_fixtures(ws):
    out = {}
    for side, (ca, cb, cc, cv, hdr) in {"thisWeek": ("A", "B", "C", "E", "A"), "nextWeek": ("H", "I", "J", "L", "H")}.items():
        wk = int(re.search(r"\(Week (\d+)\)", str(v(ws, hdr, 1))).group(1))
        out[side] = None
        for r in range(2, 34):
            away, at, home, venue = v(ws, ca, r), v(ws, cb, r), v(ws, cc, r), v(ws, cv, r)
            if isinstance(at, str) and at.startswith("BYE - ") and norm(at[6:]) == norm(TEAM):
                out[side] = {"week": wk, "bye": True}
            elif at == "@" and (is_us(away) or is_us(home)):
                us_home = is_us(home)
                out[side] = {"week": wk, "home": us_home,
                             "opponent": str(away if us_home else home).strip(),
                             "venue": str(venue).strip().strip("()")}
    return out


def page5_prediction(ws):
    lines = [str(v(ws, "A", r) or "") for r in range(36, 65)] + [str(v(ws, "H", r) or "") for r in range(35, 65)]
    text = " ".join(l.strip() for l in lines)
    text = re.sub(r"\s+", " ", text)
    for m in re.finditer(r"(Match(?:es)? #[\d #&]+:\s*\((\w) Div\.\).*?)(?=Match(?:es)? #|And the remaining|$)", text):
        if m.group(2) == DIV and TEAM.lower() in m.group(1).lower():
            return {"featured": True, "text": m.group(1).strip()}
    m = re.search(rf"{DIV} Division:\s*(.*?)(?=[A-H] Division:|$)", text.split("And the remaining", 1)[-1])
    if m:
        for part in m.group(1).split(";"):
            if TEAM.lower() in part.lower():
                return {"featured": False, "text": part.strip().rstrip(".")}
    return None


def page1_headline(ws):
    r = find_row(ws, "H", lambda x: isinstance(x, str) and x.startswith(f'"{DIV}" Division:'))
    if not r:
        return None
    parts = [str(v(ws, "H", r)).split(":", 1)[1].strip()]
    r += 1
    while isinstance(v(ws, "H", r), str) and v(ws, "H", r).startswith(" "):
        parts.append(v(ws, "H", r).strip())
        r += 1
    return re.sub(r"-\s+(?=[a-z])", "-", " ".join(parts))


def page3_hot_darts(ws):
    stop = find_row(ws, "A", lambda x: norm(x) == "divisional trophy dart categories") or ws.max_row
    out = []
    for name_c, team_c, val_c, unit_c in (("A", "B", "C", "D"), ("G", "H", "I", "J"), ("M", "N", "O", "P")):
        pending = []
        for r in range(2, stop):
            name, team = v(ws, name_c, r), v(ws, team_c, r)
            if not isinstance(name, str) or name.startswith('"') or name.startswith("("):
                pending = []
                continue
            pending.append(first_last(name))
            if team is None:
                continue  # doubles pair continues on next row
            if is_us(team):
                value, unit = v(ws, val_c, r), v(ws, unit_c, r)
                feat = f"{value} {unit}" if value not in (None, "") else str(unit)
                # 'perfect' notes sit on the row below in the unit column
                out.append({"players": pending, "feat": feat.strip()})
            pending = []
    return out


def page3_trophies(ws):
    top = find_row(ws, "A", lambda x: norm(x) == "divisional trophy dart categories")
    hdr = top + 1
    col = next(c.column_letter for c in ws[hdr] if norm(c.value) == f"{DIV.lower()} division")
    perfect = find_row(ws, "A", lambda x: norm(x) == "perfect throws", top)
    starts = [r for r in range(hdr + 1, perfect) if norm(v(ws, "A", r)) in ("high in", "high out", "fast")]
    out = []
    for i, s in enumerate(starts):
        e = starts[i + 1] if i + 1 < len(starts) else perfect
        label = " ".join(str(v(ws, "A", r)).strip() for r in range(s, e) if v(ws, "A", r))
        label = re.sub(r"\s*\(.*$", "", label)  # drop '(301 first turn only!!)' style notes
        cells = [str(v(ws, col, r)).strip() for r in range(s, e) if v(ws, col, r) not in (None, "")]
        if not cells or cells[0].startswith("(none"):
            continue
        if any(is_us(c) for c in cells):
            m = re.match(r"(.*?)\s*\((\d+)/(\d+)/(\d+)\)", cells[0])
            value, date = (m.group(1), f"20{m.group(4)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}") if m else (cells[0], None)
            players = [re.sub(r"\s*&$", "", c).strip() for c in cells[1:] if not is_us(c)]
            players = [p for c in players for p in c.split(" & ")]
            out.append({"category": label, "value": value, "date": date, "players": players})
    # perfect throws: 'Name (Team, m/d)' anywhere below
    perf = []
    for r in range(perfect, ws.max_row + 1):
        for c in ws[r]:
            if isinstance(c.value, str) and TEAM.lower() in c.value.lower():
                perf.append(c.value.strip())
    return out, perf


def leaderboard(ws, keys=("w", "l", "pct")):
    for name_c, team_c, w_c, l_c, p_c, rank_c in (("C", "D", "E", "F", "G", "A"), ("L", "M", "N", "O", "P", "J")):
        top = find_row(ws, name_c, lambda x: norm(x) == f'"{DIV.lower()}" division')
        if top:
            break
    rows, rank, r = [], None, top + 1
    while v(ws, name_c, r) not in (None, ""):
        if v(ws, rank_c, r) not in (None, ""):
            rank = v(ws, rank_c, r)
        team = str(v(ws, team_c, r) or "")
        if "/" in team and is_us(team.split("/", 1)[1]):
            rows.append({"rank": rank, "player": first_last(v(ws, name_c, r)),
                         keys[0]: num(v(ws, w_c, r)), keys[1]: num(v(ws, l_c, r)), keys[2]: round(v(ws, p_c, r), 4)})
        r += 1
    return rows


def page10(ws):
    hdr = None
    players, team_row = [], None
    for r in range(2, ws.max_row + 1):
        b = v(ws, "B", r)
        if isinstance(b, str) and re.match(rf"{DIV}\d+/ ", b):
            if team_row:
                break
            if is_us(b.split("/", 1)[1]):
                team_row = r
                code = b.split("/", 1)[0]
            continue
        if team_row and b == DIV and v(ws, "D", r):
            g = lambda c: v(ws, c, r)
            players.append({
                "number": g("C"), "name": first_last(g("D")),
                "singles": {"w": g("I"), "l": g("J")}, "doubles": {"w": g("P"), "l": g("Q")},
                "total": {"w": g("S"), "l": g("T")},
                "singles301": {"w": g("E"), "l": g("F")}, "singlesCricket": {"w": g("G"), "l": g("H")},
                "doublesCricket": {"w": g("L"), "l": g("M")}, "doubles501": {"w": g("N"), "l": g("O")},
                "tiebreaker": {"w": g("V"), "l": g("W")},
                "allStarPoints": g("X"), "gamesPlayed": g("Y"),
                "aspAverage": round(g("Z"), 4) if isinstance(g("Z"), (int, float)) else None,
                "matchesPlayed": g("AA"),
            })
    return code, players


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    season, week, pg = sheets(wb)
    date, issue = issue_info(pg("Pg1"))
    standings, last = page2(pg("Pg2"))
    fixtures = page5_fixtures(pg("Pg5"))
    code, players = page10(pg(f"Pg10{DIV}"))
    trophies, perfect = page3_trophies(pg("Pg3"))
    return {
        "season": season, "week": week, "issueDate": date, "issue": issue,
        "team": {"name": TEAM, "division": DIV, "code": code},
        "standings": standings,
        "lastResult": last,
        "thisWeek": fixtures["thisWeek"], "nextWeek": fixtures["nextWeek"],
        "prediction": page5_prediction(pg("Pg5")),
        "headline": page1_headline(pg("Pg1")),
        "players": players,
        "leaderboards": {
            "singles": leaderboard(pg("Pg6")),
            "singlesPlusDoubles": leaderboard(pg("Pg7")),
            "allStarAverage": leaderboard(pg("Pg8"), ("asp", "gamesPlayed", "average")),
        },
        "hotDarts": page3_hot_darts(pg("Pg3")),
        "trophyDarts": trophies,
        "perfectThrows": perfect,
    }


if __name__ == "__main__":
    print(json.dumps(extract(sys.argv[1]), indent=2))
